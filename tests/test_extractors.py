import importlib
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook

import pse_extract


class ExtractorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmpdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tmpdir.name)

    def tearDown(self) -> None:
        self.tmpdir.cleanup()

    def _write_text(self, name: str, content: str) -> Path:
        path = self.root / name
        path.write_text(content, encoding="utf-8")
        return path

    def _write_zip(self, name: str, members: dict[str, str]) -> Path:
        path = self.root / name
        with zipfile.ZipFile(path, "w") as archive:
            for member, content in members.items():
                archive.writestr(member, content)
        return path

    def test_html_extracts_visible_text(self) -> None:
        path = self._write_text(
            "sample.html",
            "<html><head><style>.x{}</style><script>ignore()</script></head>"
            "<body><h1>Heading</h1><p>Hello <b>world</b></p></body></html>",
        )

        text = pse_extract.extract_text(str(path), ".html", 500)

        self.assertIn("Heading", text)
        self.assertIn("Hello", text)
        self.assertNotIn("ignore()", text)

    def test_tex_strips_comments_and_commands(self) -> None:
        path = self._write_text(
            "sample.tex",
            "\\section{Intro}\nVisible text % hidden\n\\textbf{Bold words}\n\\item First item\n",
        )

        text = pse_extract.extract_text(str(path), ".tex", 500)

        self.assertIn("Intro", text)
        self.assertIn("Visible text", text)
        self.assertIn("Bold words", text)
        self.assertIn("First item", text)
        self.assertNotIn("hidden", text)

    def test_rtf_extracts_readable_text(self) -> None:
        path = self._write_text(
            "sample.rtf",
            r"{\rtf1\ansi{\fonttbl\f0\fswiss Helvetica;}\f0\pard Hello\par world\'21\par}",
        )

        text = pse_extract.extract_text(str(path), ".rtf", 500)

        self.assertIn("Hello", text)
        self.assertIn("world!", text)

    def test_docx_reads_document_xml(self) -> None:
        path = self._write_zip(
            "sample.docx",
            {
                "word/document.xml": (
                    "<?xml version='1.0' encoding='UTF-8'?>"
                    "<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
                    "<w:body><w:p><w:r><w:t>Docx body</w:t></w:r></w:p>"
                    "<w:p><w:r><w:t>Second line</w:t></w:r></w:p></w:body></w:document>"
                )
            },
        )

        text = pse_extract.extract_text(str(path), ".docx", 500)

        self.assertIn("Docx body", text)
        self.assertIn("Second line", text)

    def test_odt_reads_content_xml(self) -> None:
        path = self._write_zip(
            "sample.odt",
            {
                "content.xml": (
                    "<?xml version='1.0' encoding='UTF-8'?>"
                    "<office:document-content "
                    "xmlns:office='urn:oasis:names:tc:opendocument:xmlns:office:1.0' "
                    "xmlns:text='urn:oasis:names:tc:opendocument:xmlns:text:1.0'>"
                    "<office:body><office:text><text:h>Heading</text:h>"
                    "<text:p>Hello <text:span>ODT</text:span></text:p>"
                    "</office:text></office:body></office:document-content>"
                )
            },
        )

        text = pse_extract.extract_text(str(path), ".odt", 500)

        self.assertIn("Heading", text)
        self.assertIn("Hello ODT", text)

    def test_pptx_reads_slide_xml(self) -> None:
        path = self._write_zip(
            "sample.pptx",
            {
                "ppt/slides/slide1.xml": (
                    "<?xml version='1.0' encoding='UTF-8'?>"
                    "<p:sld xmlns:p='http://schemas.openxmlformats.org/presentationml/2006/main' "
                    "xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main'>"
                    "<p:cSld><p:spTree><p:sp><p:txBody>"
                    "<a:p><a:r><a:t>Slide title</a:t></a:r></a:p>"
                    "<a:p><a:r><a:t>Bullet text</a:t></a:r></a:p>"
                    "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
                )
            },
        )

        text = pse_extract.extract_text(str(path), ".pptx", 500)

        self.assertIn("Slide title", text)
        self.assertIn("Bullet text", text)

    def test_xlsx_reads_cell_values(self) -> None:
        path = self.root / "sample.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet["A1"] = "Hello"
        sheet["B2"] = "World"
        workbook.save(path)
        workbook.close()

        text = pse_extract.extract_text(str(path), ".xlsx", 500)

        self.assertIn("Data", text)
        self.assertIn("Hello", text)
        self.assertIn("World", text)

    def test_import_smoke_for_index_and_search_modules(self) -> None:
        pse_index = importlib.import_module("pse_index")
        pse_search = importlib.import_module("pse_search")

        self.assertTrue(callable(pse_index.index_root))
        self.assertTrue(callable(pse_search.query))

    def test_doc_com_extractor_uses_read_only_open_and_cleanup(self) -> None:
        class FakePythonCom:
            def __init__(self) -> None:
                self.calls = []

            def CoInitialize(self) -> None:
                self.calls.append("init")

            def CoUninitialize(self) -> None:
                self.calls.append("uninit")

        class FakeDocument:
            def __init__(self) -> None:
                self.closed_with = None

            def Range(self):
                return SimpleNamespace(Text="Legacy doc body")

            def Close(self, value) -> None:
                self.closed_with = value

        class FakeDocuments:
            def __init__(self, document) -> None:
                self.document = document
                self.kwargs = None

            def Open(self, path, **kwargs):
                self.kwargs = kwargs
                return self.document

        class FakeWordApp:
            def __init__(self, document) -> None:
                self.Documents = FakeDocuments(document)
                self.quit_called = False
                self.Visible = None
                self.DisplayAlerts = None

            def Quit(self) -> None:
                self.quit_called = True

        class FakeClient:
            def __init__(self, app) -> None:
                self.app = app
                self.progid = None

            def DispatchEx(self, progid: str):
                self.progid = progid
                return self.app

        pythoncom = FakePythonCom()
        document = FakeDocument()
        app = FakeWordApp(document)
        client = FakeClient(app)

        with mock.patch("pse_extract._import_pywin32", return_value=(pythoncom, client)):
            text = pse_extract._extract_doc_via_com("sample.doc", 500)

        self.assertEqual("Legacy doc body", text)
        self.assertEqual("Word.Application", client.progid)
        self.assertTrue(app.Documents.kwargs["ReadOnly"])
        self.assertFalse(document.closed_with)
        self.assertTrue(app.quit_called)
        self.assertEqual(["init", "uninit"], pythoncom.calls)

    def test_xls_com_extractor_reads_used_range_and_cleans_up(self) -> None:
        class FakePythonCom:
            def CoInitialize(self) -> None:
                return None

            def CoUninitialize(self) -> None:
                return None

        class FakeWorkbook:
            def __init__(self) -> None:
                self.closed_with = None
                self.Worksheets = [SimpleNamespace(Name="Sheet1", UsedRange=SimpleNamespace(Value=(("A1", None), ("", "B2"))))]

            def Close(self, value) -> None:
                self.closed_with = value

        class FakeWorkbooks:
            def __init__(self, workbook) -> None:
                self.workbook = workbook
                self.kwargs = None

            def Open(self, path, **kwargs):
                self.kwargs = kwargs
                return self.workbook

        class FakeExcelApp:
            def __init__(self, workbook) -> None:
                self.Workbooks = FakeWorkbooks(workbook)
                self.quit_called = False
                self.Visible = None
                self.DisplayAlerts = None

            def Quit(self) -> None:
                self.quit_called = True

        class FakeClient:
            def __init__(self, app) -> None:
                self.app = app
                self.progid = None

            def DispatchEx(self, progid: str):
                self.progid = progid
                return self.app

        workbook = FakeWorkbook()
        app = FakeExcelApp(workbook)
        client = FakeClient(app)

        with mock.patch("pse_extract._import_pywin32", return_value=(FakePythonCom(), client)):
            text = pse_extract._extract_xls_via_com("sample.xls", 500)

        self.assertIn("Sheet1", text)
        self.assertIn("A1", text)
        self.assertIn("B2", text)
        self.assertEqual("Excel.Application", client.progid)
        self.assertTrue(app.Workbooks.kwargs["ReadOnly"])
        self.assertFalse(workbook.closed_with)
        self.assertTrue(app.quit_called)

    def test_ppt_com_extractor_reads_shape_text_and_cleans_up(self) -> None:
        class FakePythonCom:
            def CoInitialize(self) -> None:
                return None

            def CoUninitialize(self) -> None:
                return None

        class FakePresentation:
            def __init__(self) -> None:
                shape = SimpleNamespace(
                    HasTextFrame=1,
                    TextFrame=SimpleNamespace(HasText=True, TextRange=SimpleNamespace(Text="Slide text")),
                )
                self.Slides = [SimpleNamespace(Shapes=[shape])]
                self.closed = False

            def Close(self) -> None:
                self.closed = True

        class FakePresentations:
            def __init__(self, presentation) -> None:
                self.presentation = presentation
                self.kwargs = None

            def Open(self, path, **kwargs):
                self.kwargs = kwargs
                return self.presentation

        class FakePowerPointApp:
            def __init__(self, presentation) -> None:
                self.Presentations = FakePresentations(presentation)
                self.quit_called = False

            def Quit(self) -> None:
                self.quit_called = True

        class FakeClient:
            def __init__(self, app) -> None:
                self.app = app
                self.progid = None

            def DispatchEx(self, progid: str):
                self.progid = progid
                return self.app

        presentation = FakePresentation()
        app = FakePowerPointApp(presentation)
        client = FakeClient(app)

        with mock.patch("pse_extract._import_pywin32", return_value=(FakePythonCom(), client)):
            text = pse_extract._extract_ppt_via_com("sample.ppt", 500)

        self.assertIn("Slide text", text)
        self.assertEqual("PowerPoint.Application", client.progid)
        self.assertTrue(app.Presentations.kwargs["ReadOnly"])
        self.assertTrue(presentation.closed)
        self.assertTrue(app.quit_called)


if __name__ == "__main__":
    unittest.main()
